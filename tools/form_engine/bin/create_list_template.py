#!/usr/bin/env python3
"""Build a small editable DOCX list template for a reviewed form pack."""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

TOOLS = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(TOOLS / "opinion_writer/bin"))


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", required=True)
    parser.add_argument("--title", required=True)
    parser.add_argument("--intro", required=True)
    parser.add_argument("--placeholder", required=True)
    parser.add_argument("--footer", default="내부 확인표 · 공식 제출 서식이 아닙니다")
    parser.add_argument("--form-font", required=True)
    parser.add_argument("--body-font", required=True)
    args = parser.parse_args(argv)

    from docx import Document
    from docx.shared import Mm, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from document_runtime import embed_fonts

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(22)
    section.left_margin = section.right_margin = Mm(24)
    styles = doc.styles
    styles["Normal"].font.name = "NanumGothic"
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    title_properties = styles["Title"].element.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run(args.title)
    run.font.name, run.font.size, run.font.bold = "NanumMyeongjo", Pt(17), True
    run.font.color.rgb = RGBColor(0, 0, 0)
    intro = doc.add_paragraph(args.intro)
    intro.paragraph_format.space_after = Pt(16)
    item = doc.add_paragraph("{{" + args.placeholder + "}}")
    item.paragraph_format.space_after = Pt(7)
    note = section.footer.paragraphs[0]
    note.text = args.footer
    note.style = styles["Normal"]
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in note.runs:
        run.font.size = Pt(8)
    output = Path(args.out)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    embed_fonts(output, {"NanumGothic": Path(args.form_font),
                         "NanumMyeongjo": Path(args.body_font)})
    print(output)


if __name__ == "__main__":
    main()
