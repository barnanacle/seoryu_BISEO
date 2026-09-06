"""Readable questionnaire documents using the same embedded, licensed fonts."""
from pathlib import Path
import json,re
from document_runtime import HERE


def write_text_document(markdown,output):
    from docx import Document
    from docx.shared import Pt,RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    template=HERE/'서식/별지_상세의견서_템플릿.docx'
    profile=json.loads((HERE/'서식/layout_profile.json').read_text())
    font=profile['fonts']['form']
    doc=Document(template)
    body=doc.element.body
    for node in list(body):
        if node.tag!=qn('w:sectPr'):body.remove(node)
    section=doc.sections[0];section.top_margin=Pt(50);section.bottom_margin=Pt(50)
    section.left_margin=Pt(50);section.right_margin=Pt(50)
    question_block=[]
    def finish_question():
        for p in question_block[:-1]:p.paragraph_format.keep_with_next=True
        for p in question_block:p.paragraph_format.keep_together=True
        question_block.clear()
    in_question=False
    for line in markdown.splitlines():
        if not line.strip():continue
        level=len(line)-len(line.lstrip('#'))
        text=line[level:].strip() if level else line
        text=text.replace('**','').replace('`','')
        if level:
            finish_question();in_question=level==3 and bool(re.match(r'Q[A-Za-z0-9_-]*\s',text))
        paragraph=doc.add_paragraph()
        if in_question:question_block.append(paragraph)
        pf=paragraph.paragraph_format;pf.first_line_indent=Pt(0);pf.left_indent=Pt(0)
        pf.space_before=Pt(10 if level else 0);pf.space_after=Pt(5)
        pf.line_spacing=Pt(21 if level else 18);pf.keep_with_next=bool(level);pf.widow_control=True
        paragraph.alignment=WD_ALIGN_PARAGRAPH.LEFT
        run=paragraph.add_run(text);run.font.name=font;run.font.size=Pt(18 if level==1 else 13 if level else 10.5)
        run.bold=bool(level);run.font.color.rgb=RGBColor(0,0,0)
        rp=run.element.get_or_add_rPr();rf=rp.find(qn('w:rFonts'))
        if rf is None:rf=OxmlElement('w:rFonts');rp.insert(0,rf)
        for name in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+name),font)
        spacing=OxmlElement('w:spacing');spacing.set(qn('w:val'),'0');rp.append(spacing)
    finish_question()
    doc.core_properties.title='행정처분 사전통지 대응 질문지';doc.core_properties.author='';doc.core_properties.last_modified_by=''
    doc.save(output)
    return Path(output)
