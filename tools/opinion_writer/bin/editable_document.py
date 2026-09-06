"""Combine an editable official-form cover and the canonical annex in one DOCX.
The blank official form is a background; the entered values are native text boxes.
"""
from __future__ import annotations
import copy,json,tempfile
from pathlib import Path
from lxml import etree

from document_runtime import HERE,render_pdf,temp_root
W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
V='urn:schemas-microsoft-com:vml'
R='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
O='urn:schemas-microsoft-com:office:office'
W10='urn:schemas-microsoft-com:office:word'


def node(parent,tag,attrs=None,text=None):
    n=etree.SubElement(parent,tag,attrs or {})
    if text is not None:n.text=text
    return n


def paragraph(parent):
    p=node(parent,'{'+W+'}p');pr=node(p,'{'+W+'}pPr')
    node(pr,'{'+W+'}spacing',{'{'+W+'}line':'20','{'+W+'}lineRule':'exact','{'+W+'}before':'0','{'+W+'}after':'0'})
    node(pr,'{'+W+'}ind',{'{'+W+'}firstLine':'0','{'+W+'}left':'0','{'+W+'}right':'0'})
    return p


def shape(p,identifier,x,y,width,height,z=1,fill=False):
    r=node(p,'{'+W+'}r');pict=node(r,'{'+W+'}pict')
    st=f'position:absolute;margin-left:{x}pt;margin-top:{y}pt;width:{width}pt;height:{height}pt;z-index:{z};mso-position-horizontal-relative:page;mso-position-vertical-relative:page'
    s=node(pict,'{'+V+'}rect',{'id':identifier,'style':st,'stroked':'f','filled':'t' if fill else 'f','fillcolor':'white'})
    node(s,'{'+W10+'}wrap',{'type':'none'})
    return s


def make_full_docx(annex_docx,blank_pdf,placements,destination,fields=None):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from pypdf import PdfReader
    profile=json.loads((HERE/'서식/layout_profile.json').read_text())
    if blank_pdf is None:return make_generic_full_docx(annex_docx,fields or {},destination)
    doc=Document(annex_docx);body=doc.element.body
    annex_section=body.find(qn('w:sectPr'))
    reset=annex_section.find(qn('w:pgNumType'))
    if reset is None:reset=OxmlElement('w:pgNumType');annex_section.append(reset)
    reset.set(qn('w:start'),'1')
    with tempfile.TemporaryDirectory(prefix='editable_cover_',dir=temp_root()) as directory:
        images=render_pdf(Path(blank_pdf),Path(directory),dpi=200)
        pages=PdfReader(blank_pdf).pages;elements=[]
        for index,(image,page) in enumerate(zip(images,pages)):
            width=float(page.mediabox.width);height=float(page.mediabox.height)
            holder=etree.Element('{'+W+'}body');p=paragraph(holder)
            rid,_=doc.part.get_or_add_image(str(image))
            bg=shape(p,f'FormBackground{index}',0,0,width,height,z=-1)
            node(bg,'{'+V+'}imagedata',{'{'+R+'}id':rid,'{'+O+'}title':'공식 빈 서식'})
            placement=placements[index] if index<len(placements) else None
            if placement:
                if placement.get('standard',False):
                    for j,(x,y,w,h) in enumerate([(398,604,137,19),(201,659,25,19)]):shape(p,f'Mask{index}_{j}',x,y,w,h,fill=True)
                for j,entry in enumerate(placement['entries']):
                    size=entry['height'];x=entry['left'];y=entry['top']+profile.get('form_docx_top_adjust_pt',-.75)
                    s=shape(p,f'Field{index}_{j}',x,y,max(entry['width']+4,10),size+8,z=2)
                    textbox=node(s,'{'+V+'}textbox',{'inset':'0,0,0,0'})
                    content=node(textbox,'{'+W+'}txbxContent');textp=paragraph(content)
                    pr=textp.find('{'+W+'}pPr');spacing=pr.find('{'+W+'}spacing');spacing.set('{'+W+'}line',str(round((size+1)*20)))
                    run=node(textp,'{'+W+'}r');rp=node(run,'{'+W+'}rPr')
                    node(rp,'{'+W+'}rFonts',{qn('w:'+key):profile['fonts']['form'] for key in ['ascii','hAnsi','eastAsia','cs']})
                    node(rp,'{'+W+'}sz',{qn('w:val'):str(round(size*2))});node(rp,'{'+W+'}spacing',{qn('w:val'):'0'})
                    node(run,'{'+W+'}t',{'{http://www.w3.org/XML/1998/namespace}space':'preserve'},entry['text'])
            section_p=paragraph(holder);sp=section_p.find('{'+W+'}pPr');section=copy.deepcopy(annex_section)
            for n in list(section):
                if n.tag in [qn('w:footerReference'),qn('w:headerReference'),qn('w:pgNumType')]:section.remove(n)
            sz=section.find(qn('w:pgSz'));sz.set(qn('w:w'),str(round(width*20)));sz.set(qn('w:h'),str(round(height*20)))
            mar=section.find(qn('w:pgMar'))
            for key in ['top','right','bottom','left','header','footer','gutter']:mar.set(qn('w:'+key),'0')
            typ=section.find(qn('w:type'))
            if typ is None:typ=OxmlElement('w:type');section.append(typ)
            typ.set(qn('w:val'),'nextPage');sp.append(section)
            elements.extend(list(holder))
        for element in reversed(elements):body.insert(0,element)
        doc.core_properties.title='의견제출서 전체 편집본';doc.save(destination)
    return Path(destination)


def make_generic_full_docx(annex_docx,fields,destination):
    from docx import Document
    from docx.shared import Pt
    from docx.enum.section import WD_SECTION_START
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from build_pdf import FORM_KEYS
    doc=Document(annex_docx);body=doc.element.body
    original_section=copy.deepcopy(body.find(qn('w:sectPr')))
    original=[copy.deepcopy(n) for n in body if n.tag!=qn('w:sectPr')]
    for n in list(body):
        if n.tag!=qn('w:sectPr'):body.remove(n)
    sec=doc.sections[0];sec.top_margin=Pt(50);sec.bottom_margin=Pt(50);sec.left_margin=Pt(50);sec.right_margin=Pt(50)
    p=doc.add_paragraph('의견제출서 초안 · 제출 서식 확인 필요')
    p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=Pt(24)
    table=doc.add_table(rows=0,cols=2)
    for key in FORM_KEYS:
        cells=table.add_row().cells;cells[0].text=key.replace('_',' ');cells[1].text=str(fields.get(key,'【확인 필요】'))
        for cell in cells:
            for p in cell.paragraphs:
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.paragraph_format.first_line_indent=Pt(0)
                p.paragraph_format.line_spacing=Pt(16);p.paragraph_format.space_after=Pt(3)
                for run in p.runs:run.font.size=Pt(9)
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    # The first section has no footer; the annex keeps its original section and page numbering.
    break_section=body.findall('.//'+qn('w:pPr')+'/'+qn('w:sectPr'))[-1]
    for n in list(break_section):
        if n.tag in [qn('w:footerReference'),qn('w:headerReference')]:break_section.remove(n)
    current=body.find(qn('w:sectPr'));body.remove(current)
    for n in original:body.append(n)
    pg=original_section.find(qn('w:pgNumType'))
    if pg is None:pg=etree.SubElement(original_section,qn('w:pgNumType'))
    pg.set(qn('w:start'),'1');body.append(original_section)
    doc.save(destination);return Path(destination)


def validate_full_docx(full_docx,annex_pdf,placements,work_dir):
    from document_runtime import convert_docx
    from validate_layout import compare_images,compact,find_phrase
    from pypdf import PdfReader
    import pdfplumber
    work_dir=Path(work_dir);full_pdf=work_dir/'full_editor.pdf'
    convert_docx(full_docx,full_pdf)
    annex_pages=len(PdfReader(annex_pdf).pages);cover_count=len(PdfReader(full_pdf).pages)-annex_pages
    errors=[]
    if cover_count<1:errors.append('전체 편집본의 표지가 없습니다.')
    if placements and cover_count!=len(placements):errors.append('전체 편집본의 표지 쪽수가 다릅니다.')
    with pdfplumber.open(full_pdf) as pdf:
        for index,placement in enumerate(placements):
            if not placement:continue
            page=pdf.pages[index];text=compact(page.extract_text() or '')
            for entry in placement['entries']:
                if compact(entry['text']) not in text:errors.append('전체 편집본 표지 값 누락: '+entry['field'])
                found=find_phrase(page,entry['text'],entry['top']-2,entry['top']+2,min_x=entry['left']-2,max_x=entry['left']+entry['width']+2)
                if found is None:errors.append('전체 편집본 표지 위치 오류: '+entry['field'])
    actual=render_pdf(full_pdf,work_dir/'full_check',first=cover_count+1)
    expected=render_pdf(Path(annex_pdf),work_dir/'annex_check')
    comparisons=[compare_images(a,b) for a,b in zip(actual,expected)]
    if len(actual)!=len(expected) or any(not c['equal'] for c in comparisons):errors.append('전체 편집본의 별지가 원본 DOCX 별지와 다릅니다.')
    return {'passed':not errors,'errors':errors,'cover_pages':cover_count,'annex_pages':annex_pages,'comparisons':comparisons}
