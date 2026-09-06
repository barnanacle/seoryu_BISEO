"""Rebuild the approved DOCX layout template; run only for an explicit layout change."""
from pathlib import Path
import json
from document_runtime import bootstrap


def create_template(destination: Path):
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from document_runtime import HERE, embed_fonts, extract_fonts, temp_root
    import tempfile

    profile = json.loads((HERE / '서식/layout_profile.json').read_text())
    spec = profile['annex']
    font = profile['fonts']['body']
    font_temp=None
    if profile.get('font_files'):
        font_paths={profile['fonts'][kind]:HERE/path for kind,path in profile['font_files'].items()}
    else:
        font_temp=tempfile.TemporaryDirectory(prefix='template_fonts_',dir=temp_root())
        font_paths=extract_fonts(destination,Path(font_temp.name))
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = map(Pt, profile['page_pt'])
    sec.left_margin = Pt(spec['left_pt']); sec.right_margin = Pt(spec['right_pt'])
    sec.top_margin = Pt(spec['top_pt']); sec.bottom_margin = Pt(spec['bottom_pt'])
    sec.header_distance = Pt(24); sec.footer_distance = Pt(34)
    for node in list(sec._sectPr):
        if node.tag == qn('w:docGrid'):
            sec._sectPr.remove(node)

    def set_font(target, size=12, bold=False, tracking=0):
        target.font.name = font; target.font.size = Pt(size)
        target.font.bold = bold; target.font.color.rgb = RGBColor(0,0,0)
        rp = target.element.get_or_add_rPr()
        rf = rp.find(qn('w:rFonts'))
        if rf is None:
            rf = OxmlElement('w:rFonts'); rp.insert(0, rf)
        for key in ['ascii','hAnsi','eastAsia','cs']:
            rf.set(qn('w:'+key), font)
        for key in list(rf.attrib):
            if key.endswith('Theme'):
                del rf.attrib[key]
        space = rp.find(qn('w:spacing'))
        if space is None:
            space = OxmlElement('w:spacing'); rp.append(space)
        space.set(qn('w:val'), str(round(tracking*20)))
        lang=OxmlElement('w:lang');lang.set(qn('w:val'),'ko-KR');lang.set(qn('w:eastAsia'),'ko-KR');rp.append(lang)

    normal = doc.styles['Normal']
    set_font(normal, spec['body_size_pt'], tracking=spec['body_tracking_pt'])
    normal.paragraph_format.line_spacing=Pt(spec['line_pt'])
    normal.paragraph_format.space_before=Pt(0);normal.paragraph_format.space_after=Pt(0)
    normal.paragraph_format.first_line_indent=Pt(spec['first_indent_pt'])
    for name in ['Title','Heading 1','Heading 2']:
        st=doc.styles[name];set_font(st, 20 if name=='Title' else 15, True, 0)
        for b in st.element.xpath('./w:pPr/w:pBdr'):
            b.getparent().remove(b)

    def fmt(p, *, size=12, bold=False, tracking=0, line=30, before=0, after=0, indent=0, center=False, keep=False):
        pf=p.paragraph_format
        pf.line_spacing=Pt(line);pf.space_before=Pt(before);pf.space_after=Pt(after)
        pf.first_line_indent=Pt(indent);pf.left_indent=Pt(0);pf.right_indent=Pt(0)
        pf.keep_with_next=keep;pf.widow_control=True
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
        pp=p._p.get_or_add_pPr()
        for key in ['snapToGrid','autoSpaceDE','autoSpaceDN','adjustRightInd']:
            node=OxmlElement('w:'+key);node.set(qn('w:val'),'0');pp.append(node)
        for run in p.runs:set_font(run,size,bold,tracking)
        return p

    fmt(doc.add_paragraph('[별지]'),keep=True).alignment=WD_ALIGN_PARAGRAPH.LEFT
    title=doc.add_paragraph('{{별지_제목}}',style='Title')
    fmt(title,size=spec['title_size_pt'],bold=True,line=spec['title_line_pt'],before=spec['title_before_pt'],after=spec['title_after_pt'],center=True,keep=True)

    table=doc.add_table(rows=3,cols=3)
    table.alignment=WD_TABLE_ALIGNMENT.LEFT;table.autofit=False
    total_width=595-spec['left_pt']-spec['right_pt']
    widths=[spec['label_width_pt'],spec['label_gap_pt'],total_width-spec['label_width_pt']-spec['label_gap_pt']]
    tp=table._tbl.tblPr
    # Word derives an auto table width from its grid; HTML-based viewers can shrink it.
    table_width=tp.find(qn('w:tblW'))
    table_width.set(qn('w:type'),'dxa');table_width.set(qn('w:w'),str(round(total_width*20)))
    for tag in ['tblBorders','tblCellMar','tblInd']:
        old=tp.find(qn('w:'+tag))
        if old is not None:tp.remove(old)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        child=OxmlElement('w:'+edge);child.set(qn('w:val'),'nil');borders.append(child)
    tp.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for edge in ['top','left','bottom','right']:
        child=OxmlElement('w:'+edge);child.set(qn('w:w'),'0');child.set(qn('w:type'),'dxa');margins.append(child)
    tp.append(margins)
    indent=OxmlElement('w:tblInd');indent.set(qn('w:w'),'0');indent.set(qn('w:type'),'dxa');tp.append(indent)
    for i,w in enumerate(widths):table.columns[i].width=Pt(w)
    metadata=[('의견제출인','{{의견제출인_명칭}}'),('','{{별지_주소}}'),('처\t분\t청','{{처분청}}')]
    for row,(label,value) in zip(table.rows,metadata):
        row.height=Pt(30);row.height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for cell,width in zip(row.cells,widths):
            cell.width=Pt(width)
            cell_margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                node=OxmlElement('w:'+edge);node.set(qn('w:w'),'0');node.set(qn('w:type'),'dxa');cell_margins.append(node)
            cell._tc.get_or_add_tcPr().append(cell_margins)
        for cell,text in [(row.cells[0],label),(row.cells[1],''),(row.cells[2],value)]:
            cell.text=text;p=fmt(cell.paragraphs[0],keep=True)
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        row.cells[0].paragraphs[0].paragraph_format.tab_stops.add_tab_stop(Pt(24))
        row.cells[0].paragraphs[0].paragraph_format.tab_stops.add_tab_stop(Pt(48))
        if value=='{{별지_주소}}':
            for run in row.cells[2].paragraphs[0].runs:
                scale=OxmlElement('w:w');scale.set(qn('w:val'),str(spec['address_scale_percent']));run.element.get_or_add_rPr().append(scale)

    fmt(doc.add_paragraph('{{건명}}'),bold=True,before=30,keep=True).alignment=WD_ALIGN_PARAGRAPH.LEFT
    target=fmt(doc.add_paragraph('{{대상_표시}}'),bold=True,tracking=spec['target_tracking_pt'],keep=True)
    target.alignment=WD_ALIGN_PARAGRAPH.LEFT
    for run in target.runs:
        scale=OxmlElement('w:w');scale.set(qn('w:val'),str(spec['target_scale_percent']));run.element.get_or_add_rPr().append(scale)
    fmt(doc.add_paragraph('의 견 제 출 의  취 지',style='Heading 1'),size=15,bold=True,line=37.5,before=spec['intent_before_pt'],after=spec['heading_after_pt'],center=True,keep=True)
    fmt(doc.add_paragraph('{{취지}}'),tracking=spec['body_tracking_pt'],indent=66)
    fmt(doc.add_paragraph('사실관계 및 의견의 근거',style='Heading 1'),size=15,bold=True,line=37.5,before=spec['section_before_pt'],after=spec['heading_after_pt'],center=True,keep=True)
    fmt(doc.add_paragraph('{{소명_제목}}',style='Heading 2'),size=12,bold=True,before=15,keep=True).alignment=WD_ALIGN_PARAGRAPH.LEFT
    fmt(doc.add_paragraph('{{소명_본문}}'),tracking=spec['body_tracking_pt'],indent=66)
    fmt(doc.add_paragraph('결    어',style='Heading 1'),size=15,bold=True,line=37.5,before=spec['section_before_pt'],after=spec['heading_after_pt'],center=True,keep=True)
    fmt(doc.add_paragraph('{{결어}}'),tracking=spec['body_tracking_pt'],indent=66)
    fmt(doc.add_paragraph('{{자료목록_제목}}',style='Heading 1'),size=15,bold=True,line=37.5,before=45,after=15,center=True,keep=True)
    fmt(doc.add_paragraph('{{소명자료_목록}}'),tracking=0).alignment=WD_ALIGN_PARAGRAPH.LEFT
    fmt(doc.add_paragraph('{{별지_작성일}}'),before=36,center=True,keep=True)
    fmt(doc.add_paragraph('위 의견제출인   {{별지_서명}}'),center=True)
    footer=sec.footer.paragraphs[0]
    footer.add_run('-  ')
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE')
    r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='1';r.append(t);fld.append(r);footer._p.append(fld)
    footer.add_run('  -')
    fmt(footer,size=10,line=12,center=True)
    settings=doc.settings.element
    update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
    doc.core_properties.title='처분사전통지 상세 의견서 서식'
    doc.core_properties.author='';doc.core_properties.last_modified_by=''
    doc.save(destination)
    embed_fonts(destination,font_paths)
    if font_temp:font_temp.cleanup()


if __name__=='__main__':
    bootstrap()
    from document_runtime import HERE
    create_template(HERE/'서식/별지_상세의견서_템플릿.docx')
    print('제출본 기준 DOCX 템플릿 생성 완료')
